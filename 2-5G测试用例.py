import re

from openpyxl import load_workbook

wb = load_workbook('5GNRChannel.xlsx')
print(len(wb.sheetnames))
# 获取到所有表的名字,是一个列表
print(wb.sheetnames)
# 通过表名字获取指定表
sheet_name = "N15_15kHz"
sheet = wb[sheet_name]
# 输出表的最大行和最大列
print(sheet.max_row)
print(sheet.max_column)
# 读取表中的数据Bandwidth [MHz],Carrier centre,
# 读出每一行所有值
col_dict = {'1': 'A', '2': 'B', '3': 'C', '4': 'D', '5': 'E', '6': 'F', '7': 'G', '8': 'H', '9': 'I', '10': 'J',
            '11': 'K', '12': 'L', '13': 'M', '14': 'N', '15': 'O', '16': 'P', '17': 'Q', '18': 'R'}
word_allow_list = ['Bandwidth [MHz]', 'Carrier centre', 'absoluteFrequencyPointA', 'absoluteFrequencySSB']
word_value_list = []  # 第一行数据的列和值
for i in sheet['1']:
    print(i.row, i.column, i.value)
    if i.value in word_allow_list:
        word_value_list.append((col_dict[str(i.column)], i.value))
print(word_value_list)
# 获取指定列中数字的行和值列表
num_value_list = []
print(22222222222222)
for col, val in word_value_list:
    if val in ['Bandwidth [MHz]']:
        for x in sheet[col]:
            if isinstance(x.value, int):
                num_value_list.append((str(x.row), x.value))
print(num_value_list)
print(33333333333333)

# 获取up系列的值，先判断是否有符号&，符号实在第三列C中,先获取到C列值，判断&是否在其中,在其中的就从之前的Down中获取up值
c_value_list = []
for i in sheet['C']:
    c_value_list.append(i.value)
# 符号不在其中的话，说明有up列的数据，所以需要再获取up的数据，
if "&" not in c_value_list:
    print("不在其中")
else:
    # 否则的话说明&在，就只需要从down中拼接up的数据
    pass
# 找到需要的数据值
sheet_value_dict = dict()
# 遍历行,获取到Down系列的值
for row, val in num_value_list:  # 里面每个元素表示数字一系列
    # 遍历列
    data_value_list = []
    for col, word in word_value_list:  # 每个元素表示一列
        # 符号不在其中的话，说明有up列的数据，所以需要再获取up的数据，
        if "&" not in c_value_list:
            print("&&&&&&&&&&不在其中")
            # 遍历行实现加一行操作
            print(sheet[col + row].value)
            print(sheet[col + str(int(row) + 1)].value)
            print(sheet[col + str(int(row) + 2)].value)
            print(sheet[col + str(int(row) + 3)].value)
            print(sheet[col + str(int(row) + 4)].value)
            print(sheet[col + str(int(row) + 5)].value)
            data_value_list += [sheet[col + row].value, sheet[col + str(int(row) + 1)].value,
                                sheet[col + str(int(row) + 2)].value, sheet[col + str(int(row) + 3)].value,
                                sheet[col + str(int(row) + 4)].value, sheet[col + str(int(row) + 5)].value]
        else:
            # 否则的话说明&在，就只需要从down中拼接up的数据
            print("&&&&&&&&&&在其中")
            # 遍历行实现加一行操作
            print(sheet[col + row].value)
            print(sheet[col + str(int(row) + 1)].value)
            print(sheet[col + str(int(row) + 2)].value)
            data_value_list += [sheet[col + row].value, sheet[col + str(int(row) + 1)].value,
                                sheet[col + str(int(row) + 2)].value, sheet[col + row].value,
                                sheet[col + str(int(row) + 1)].value, sheet[col + str(int(row) + 2)].value, ]

    print(4444444444444444444444444444444)
    sheet_value_dict[data_value_list[0]] = data_value_list
print(sheet_value_dict)
# 拼接字符串到目标字符串中
start = re.match(r'N(\d*)_(\d*)kHz', sheet_name)
title_str = f"""String [][] n{start.group(1)} = new String[][] (
                // n{start.group(1)}-{start.group(2)}"""
for key, value in sheet_value_dict.items():
    str_template = f"""
                    ("{key}", "{start.group(2)}", "LOW", "{value[6]}", "{value[12]}", "{value[18]}", "{value[24]}", "{value[9]}", "{value[15]}",  "{value[21]}", "0" ),
                    ("{key}", "{start.group(2)}", "MID", "{value[7]}", "{value[13]}", "{value[19]}", "{value[25]}", "{value[10]}", "{value[16]}",  "{value[22]}", "0" ),
                    ("{key}", "{start.group(2)}", "HIGH", "{value[8]}", "{value[14]}", "{value[20]}", "{value[26]}", "{value[11]}", "{value[17]}",  "{value[23]}", "0" ),
                    """
    print(str_template)
    print(55555555555555555555555555555555)
